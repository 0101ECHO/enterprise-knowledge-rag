"""
文档解析器 - 各格式文档的具体解析实现
PDF: pdfplumber (保留表格结构) + OCR
Word: python-docx
HTML/网页: BeautifulSoup
图片: RapidOCR
"""

import os
from typing import Optional

from app.core.config import settings
from app.core.logging import log
from app.services.document.loader import ParsedDocument


def parse_pdf(file_path: str) -> ParsedDocument:
    """解析 PDF 文档 - 保留表格结构与图片 OCR"""
    import pdfplumber

    doc = ParsedDocument(title=os.path.basename(file_path))
    pages = []

    with pdfplumber.open(file_path) as pdf:
        doc.metadata["page_count"] = len(pdf.pages)
        doc.metadata["author"] = pdf.metadata.get("Author", "") if pdf.metadata else ""

        for i, page in enumerate(pdf.pages, 1):
            page_text = page.extract_text() or ""

            # 提取表格
            tables = []
            try:
                extracted_tables = page.extract_tables()
                for table in extracted_tables:
                    if table and len(table) > 1:
                        # 将表格转为 Markdown 格式
                        md_table = _table_to_markdown(table)
                        tables.append(md_table)
                        page_text += f"\n\n[表格]\n{md_table}\n"
            except Exception as e:
                log.warning(f"PDF 表格提取失败 (page {i}): {e}")

            # 提取图片并进行 OCR
            images = []
            if settings.OCR_ENABLED:
                try:
                    for img in page.images:
                        try:
                            # 裁剪图片区域
                            bbox = (
                                img["x0"],
                                page.height - img["y1"],
                                img["x1"],
                                page.height - img["y0"],
                            )
                            cropped = page.within_bbox(bbox).to_image(
                                resolution=200
                            )
                            with _temp_image() as tmp_path:
                                cropped.save(tmp_path)
                                ocr_text = _ocr_image(tmp_path)
                                if ocr_text.strip():
                                    images.append(ocr_text)
                                    page_text += f"\n[图片OCR]\n{ocr_text}\n"
                        except Exception as e:
                            log.debug(f"图片 OCR 失败 (page {i}): {e}")
                except Exception as e:
                    log.warning(f"PDF 图片提取失败 (page {i}): {e}")

            pages.append({
                "page": i,
                "text": page_text.strip(),
                "tables": tables,
                "images": images,
            })

    doc.pages = pages
    doc.raw_text = "\n\n".join(p["text"] for p in pages)
    return doc


def parse_docx(file_path: str) -> ParsedDocument:
    """解析 Word 文档"""
    from docx import Document as DocxDocument

    doc = ParsedDocument(title=os.path.basename(file_path))
    docx = DocxDocument(file_path)

    # 提取元数据
    core_props = docx.core_properties
    doc.metadata["author"] = core_props.author or ""
    doc.metadata["title"] = core_props.title or ""
    doc.metadata["created"] = str(core_props.created) if core_props.created else ""

    full_text_parts = []
    tables_text = []

    # 遍历文档体，保持段落和表格的顺序
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from docx.oxml.ns import qn

    body = docx.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            para = Paragraph(child, docx)
            if para.text.strip():
                full_text_parts.append(para.text)
        elif child.tag == qn("w:tbl"):
            table = Table(child, docx)
            table_data = []
            for row in table.rows:
                row_data = [cell.text.strip() for cell in row.cells]
                table_data.append(row_data)
            if table_data:
                md_table = _table_to_markdown(table_data)
                tables_text.append(md_table)
                full_text_parts.append(f"\n[表格]\n{md_table}\n")

    # 提取内嵌图片 OCR
    if settings.OCR_ENABLED:
        try:
            from docx.opc.constants import RELATIONSHIP_TYPE as RT
            for rel in docx.part.rels.values():
                if "image" in rel.reltype:
                    try:
                        image_data = rel.target_part.blob
                        with _temp_image_bytes(image_data) as tmp_path:
                            ocr_text = _ocr_image(tmp_path)
                            if ocr_text.strip():
                                full_text_parts.append(f"\n[图片OCR]\n{ocr_text}\n")
                    except Exception as e:
                        log.debug(f"Word 图片 OCR 失败: {e}")
        except Exception as e:
            log.warning(f"Word 图片提取失败: {e}")

    text = "\n".join(full_text_parts)
    doc.pages = [{"page": 1, "text": text, "tables": tables_text, "images": []}]
    doc.raw_text = text
    doc.metadata["table_count"] = len(tables_text)
    return doc


def parse_html(file_path: str) -> ParsedDocument:
    """解析 HTML / 网页文档"""
    from bs4 import BeautifulSoup

    doc = ParsedDocument(title=os.path.basename(file_path))

    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        html_content = f.read()

    soup = BeautifulSoup(html_content, "lxml")

    # 提取标题
    title_tag = soup.find("title")
    if title_tag:
        doc.title = title_tag.get_text(strip=True)

    # 移除 script/style 标签
    for tag in soup(["script", "style", "nav", "footer", "header"]):
        tag.decompose()

    # 提取表格
    tables = []
    for table in soup.find_all("table"):
        table_data = []
        for row in table.find_all("tr"):
            cells = row.find_all(["td", "th"])
            row_data = [cell.get_text(strip=True) for cell in cells]
            if row_data:
                table_data.append(row_data)
        if table_data:
            md_table = _table_to_markdown(table_data)
            tables.append(md_table)

    # 提取正文文本
    text = soup.get_text(separator="\n", strip=True)
    # 合并多余空行
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    text = "\n".join(lines)

    doc.pages = [{"page": 1, "text": text, "tables": tables, "images": []}]
    doc.raw_text = text
    doc.metadata["table_count"] = len(tables)
    return doc


def parse_markdown(file_path: str) -> ParsedDocument:
    """解析 Markdown 文档"""
    doc = ParsedDocument(title=os.path.basename(file_path))

    with open(file_path, "r", encoding="utf-8") as f:
        text = f.read()

    # 提取标题 (第一个 # 标题)
    for line in text.splitlines():
        if line.startswith("# "):
            doc.title = line[2:].strip()
            break

    doc.pages = [{"page": 1, "text": text, "tables": [], "images": []}]
    doc.raw_text = text
    return doc


def parse_text(file_path: str) -> ParsedDocument:
    """解析纯文本文档"""
    doc = ParsedDocument(title=os.path.basename(file_path))

    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        text = f.read()

    doc.pages = [{"page": 1, "text": text, "tables": [], "images": []}]
    doc.raw_text = text
    return doc


def parse_image(file_path: str) -> ParsedDocument:
    """解析图片 - 使用 OCR 提取文本"""
    doc = ParsedDocument(title=os.path.basename(file_path))

    ocr_text = _ocr_image(file_path)
    doc.pages = [{"page": 1, "text": ocr_text, "tables": [], "images": [ocr_text]}]
    doc.raw_text = ocr_text
    doc.metadata["ocr"] = True
    return doc


# ========== 辅助函数 ==========

def _table_to_markdown(table_data: list[list[str]]) -> str:
    """将二维表格数据转为 Markdown 格式"""
    if not table_data:
        return ""

    # 清理单元格内容
    cleaned = []
    for row in table_data:
        cleaned_row = [str(cell).replace("\n", " ").replace("|", "\\|").strip() for cell in row]
        cleaned.append(cleaned_row)

    # 构建 Markdown 表格
    header = cleaned[0]
    separator = ["---"] * len(header)
    body = cleaned[1:] if len(cleaned) > 1 else []

    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(separator) + " |",
    ]
    for row in body:
        # 补齐列数
        while len(row) < len(header):
            row.append("")
        lines.append("| " + " | ".join(row) + " |")

    return "\n".join(lines)


def _ocr_image(image_path: str) -> str:
    """对图片执行 OCR"""
    if not settings.OCR_ENABLED:
        return ""

    try:
        from rapidocr_onnxruntime import RapidOCR

        ocr = RapidOCR()
        result, _ = ocr(image_path)
        if result:
            texts = [item[1] for item in result]
            return "\n".join(texts)
    except ImportError:
        log.warning("RapidOCR 未安装，跳过 OCR")
    except Exception as e:
        log.warning(f"OCR 失败: {e}")

    return ""


import contextlib


@contextlib.contextmanager
def _temp_image(suffix: str = ".png"):
    """临时图片文件上下文管理器"""
    import tempfile

    tmp = tempfile.NamedTemporaryFile(suffix=suffix, delete=False)
    tmp.close()
    try:
        yield tmp.name
    finally:
        if os.path.exists(tmp.name):
            os.unlink(tmp.name)


@contextlib.contextmanager
def _temp_image_bytes(data: bytes, suffix: str = ".png"):
    """临时图片文件(字节内容)上下文管理器"""
    import tempfile

    tmp = tempfile.NamedTemporaryFile(suffix=suffix, delete=False)
    tmp.write(data)
    tmp.close()
    try:
        yield tmp.name
    finally:
        if os.path.exists(tmp.name):
            os.unlink(tmp.name)
